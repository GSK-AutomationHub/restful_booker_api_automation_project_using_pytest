import requests
import openpyxl
from openpyxl.styles import PatternFill

""" this utility consist function to read write from / to excel file """


def get_active_sheet(file, sheetName):
    workbook = openpyxl.load_workbook(file)
    active_sheet = workbook[sheetName]
    return active_sheet


def get_max_row(sheetName):
    return sheetName.max_row


def get_max_column(sheetName):
    return sheetName.max_column


def read_data_from_excel(file, sheetName, rownum, colnum):
    activeSheet = get_active_sheet(file, sheetName)
    print(activeSheet.cell(rownum, colnum).value, end='            ')
    return activeSheet.cell(rownum, colnum).value


def read_all_data(file, sheetName):
    activeSheet = get_active_sheet(file, sheetName)
    rows = get_max_row(activeSheet)
    columns = get_max_column(activeSheet)

    for r in range(1, rows + 1):
        for c in range(1, columns + 1):
            print(activeSheet.cell(r, c).value, end='            ')
        print()


def read_save_all_data(file, sheetName):
    excelData = []
    activeSheet = get_active_sheet(file, sheetName)
    rows = get_max_row(activeSheet)
    columns = get_max_column(activeSheet)

    for r in range(2, rows + 1):
        for c in range(1, columns + 1):
            if c < 7:
                excelData.append(activeSheet.cell(r, c).value)

    return excelData
    # print(len(excelData))
    # for data in excelData:
    #     print(data)


def fill_green(file, sheetName, rownum, colnum):
    workbook = openpyxl.load_workbook(file)
    activeSheet = workbook[sheetName]
    GreenFill = PatternFill(start_color='60b212', end_color='60b212', fill_type='solid')
    activeSheet.cell(rownum, colnum).fill = GreenFill
    workbook.save(file)


def fill_red(file, sheetName, rownum, colnum):
    workbook = openpyxl.load_workbook(file)
    activeSheet = workbook[sheetName]
    RedFill = PatternFill(start_color='ff0000', end_color='ff0000', fill_type='solid')
    activeSheet.cell(rownum, colnum).fill = RedFill
    workbook.save(file)
